"""Helpers for filling templates/Shift-Projects-Template.pptx with python-pptx.

    from shift_pptx import open_template, keep, duplicate, set_text, find, texts
    prs = open_template()
    keep(prs, [1, 4, 8, 19])            # 1-based slide numbers to keep, in this order
    s = prs.slides[1]
    set_text(s, 'Heading', 'Built for today,\\nready for tomorrow')
    prs.save('deck.pptx')               # then drag into Google Drive -> open with Google Slides

Every shape on a slide is named (Title, Heading, Stat 1 number, Card 2 text ...). Use texts(slide) to list them.
"""
import copy, os
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, '..', 'templates', 'Shift-Projects-Template.pptx')


def open_template(path=TEMPLATE):
    return Presentation(path)


def keep(prs, numbers):
    """Keep only the given 1-based slides, in the given order (a number may repeat to duplicate that slide)."""
    originals = list(prs.slides)
    ids = list(prs.slides._sldIdLst)
    wanted = [originals[n - 1] for n in numbers]
    # duplicates first (python-pptx cannot reorder, so rebuild the list at the end)
    new_slides = []
    used = set()
    for s in wanted:
        if id(s) in used:
            new_slides.append(duplicate(prs, s))
        else:
            used.add(id(s)); new_slides.append(s)
    for sid, s in zip(ids, originals):
        if id(s) not in used:
            prs.part.drop_rel(sid.rId); prs.slides._sldIdLst.remove(sid)
    # reorder sldIdLst to match requested order
    lst = prs.slides._sldIdLst
    order = [next(sid for sid in list(lst) if prs.part.related_part(sid.rId).slide is s) for s in new_slides]
    for sid in list(lst): lst.remove(sid)
    for sid in order: lst.append(sid)
    return prs


def duplicate(prs, slide):
    """Append a copy of `slide` (all template slides are native shapes, so a shape-tree copy is exact)."""
    new = prs.slides.add_slide(slide.slide_layout)
    for shp in list(new.shapes): shp._element.getparent().remove(shp._element)
    for shp in slide.shapes:
        new.shapes._spTree.append(copy.deepcopy(shp._element))
    if slide.has_notes_slide:
        new.notes_slide.notes_text_frame.text = slide.notes_slide.notes_text_frame.text
    return new


def find(slide, name):
    return next((s for s in slide.shapes if s.name == name), None)


def texts(slide):
    """{shape name: text} for every text shape on the slide."""
    return {s.name: s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()}


def set_text(slide, name, value):
    """Replace a named shape's text, keeping the first run's formatting. '\\n' makes new paragraphs."""
    shp = find(slide, name)
    if shp is None:
        raise KeyError(f'no shape named {name!r}; have: {sorted(texts(slide))}')
    tf = shp.text_frame
    lines = value.split('\n')
    p0 = tf.paragraphs[0]
    run0 = p0.runs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    run0.text = lines[0]
    for ln in lines[1:]:
        newp = copy.deepcopy(p0._p)
        p0._p.getparent().append(newp)
        from pptx.text.text import _Paragraph
        _Paragraph(newp, tf).runs[0].text = ln
    return shp
